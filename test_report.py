from pathlib import Path
import openpyxl
import pytest
from report import load_sales, analyze, write_report


def test_sample_aggregates():
    frame=load_sales("data/demo_sales.csv")
    assert int(frame["gross_cents"].sum())==209750
    assert int(frame["quantity"].sum())==21
    summary=analyze(frame)
    assert summary.iloc[0]["product"]=="Filtre kahve / Filter coffee"


@pytest.mark.parametrize("row",["X,2026-10-01,Latte,-1,10", "X,2026-10-01,Latte,1,NaN", "X,2026-10-01,Latte,1,1.234", "X,wrong,Latte,1,10", "X,2026-10-01,,1,10"])
def test_bad_input(tmp_path,row):
    source=tmp_path/"bad.csv";source.write_text("order_id,date,product,quantity,unit_price\n"+row,encoding="utf-8")
    with pytest.raises(ValueError):load_sales(source)


def test_duplicate_ids(tmp_path):
    source=tmp_path/"dupe.csv";source.write_text("order_id,date,product,quantity,unit_price\nX,2026-10-01,Latte,1,10\nX,2026-10-01,Latte,2,10",encoding="utf-8")
    with pytest.raises(ValueError,match="Duplicate"):load_sales(source)


@pytest.mark.parametrize("language,total_sheet,raw_sheet",[("tr","Özet","Satışlar"),("en","Summary","Sales")])
def test_excel_formulas_cached_totals_and_chart(tmp_path,language,total_sheet,raw_sheet):
    output=tmp_path/f"{language}.xlsx";write_report(load_sales("data/demo_sales.csv"),output,language)
    formulas=openpyxl.load_workbook(output,data_only=False)
    values=openpyxl.load_workbook(output,data_only=True)
    assert formulas[total_sheet]["B4"].value.startswith("=SUM(")
    assert formulas[raw_sheet]["F6"].value=="=D6*E6"
    assert values[total_sheet]["B4"].value==2097.5
    assert values[total_sheet]["B5"].value==21
    assert len(formulas[total_sheet]._charts)==1
    assert values[raw_sheet]["B6"].is_date
    assert sum(values[total_sheet].cell(row,3).value for row in range(8,11))==2097.5


def test_spreadsheet_formula_input_remains_text(tmp_path):
    source=tmp_path/"text.csv";source.write_text("order_id,date,product,quantity,unit_price\nX,2026-10-01,=1+1,1,0",encoding="utf-8")
    output=tmp_path/"safe.xlsx";write_report(load_sales(source),output,"en")
    workbook=openpyxl.load_workbook(output)
    assert workbook["Sales"]["C6"].data_type=="s"
